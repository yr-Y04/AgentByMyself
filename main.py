import json
import os
import re
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional
from urllib.parse import quote, unquote, urlparse

import requests
from bs4 import BeautifulSoup
from docx import Document
from dotenv import load_dotenv
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.edge.options import Options
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

load_dotenv()

# ====================== 可直接修改：系统提示词 ======================
LLM_SYSTEM_PROMPT = """
你是一个“文件工作助手”，服务对象是不太会使用电脑和 AI 的非计算机专业中老年人。

你的回答必须遵守下面要求：
1. 一定要耐心。
2. 解释要详细、通俗易懂，尽量少用专业术语。
3. 如果必须用术语，要顺便解释术语是什么意思。
4. 操作步骤要分步骤说明，像教人一步一步做事一样。
5. 如果用户要求你把内容“写入文件”，你必须优先调用写文件工具，而不是只在聊天里输出内容。
6. 如果用户要求“参考某个资料文件”来写作，必须先读取资料文件，再基于读取结果生成内容。
7. 如果用户要求联网查找资料，优先使用 web_search_bing。
8. 如果用户要求搜索最近新闻，优先使用 search_news_bing。
9. 如果用户要求搜索图片，优先使用 image_search_bing。
10. 如果用户要求下载图片，优先使用 download_bing_image，而不是先让自己编造图片链接。
11. 如果用户要求把网上图片、文档或其他文件保存到本地，使用 download_url。
12. 如果用户询问今天日期、当前时间、现在几点，优先使用 get_current_datetime。
13. 如果一个任务需要多个步骤，请连续调用多个工具，直到任务完成。
14. 如果工具已经成功执行，不要再说“我无法写入文件”这类与事实不符的话。
15. 不要编造已经读取到的文件内容；凡是涉及文件内容，都以工具返回结果为准。
16. 当新闻搜索工具已经返回结果时，你必须依据工具结果直接总结，不要再说“我无法提供最新新闻”之类的话。
17. 如果失败，请明确说明失败原因，并提出下一步建议。
""".strip()
# =================================================================

# ====================== 智谱 API 配置 ======================
LLM_API_URL = "https://open.bigmodel.cn/api/paas/v4/chat/completions"
LLM_MODEL = os.getenv("ZHIPU_MODEL", "glm-4-flash-250414")
LLM_API_KEY = os.getenv("ZHIPU_API_KEY")
# =========================================================

WORKSPACE_ROOT = Path(os.getenv("AGENT_WORKSPACE", "./agent_workspace")).resolve()

TEXT_EXTENSIONS = {
    ".txt", ".md", ".json", ".csv", ".py", ".java", ".js", ".ts",
    ".html", ".css", ".xml", ".yaml", ".yml", ".log"
}

MAX_AGENT_STEPS = 8
MAX_RECENT_ROUNDS = 6

MEMORY_SUMMARY_FILE = WORKSPACE_ROOT / "_memory_summary.txt"
RECENT_HISTORY_FILE = WORKSPACE_ROOT / "_recent_history.txt"


def ensure_workspace() -> None:
    WORKSPACE_ROOT.mkdir(parents=True, exist_ok=True)
    (WORKSPACE_ROOT / "downloads").mkdir(parents=True, exist_ok=True)

    if not MEMORY_SUMMARY_FILE.exists():
        MEMORY_SUMMARY_FILE.write_text("", encoding="utf-8")

    if not RECENT_HISTORY_FILE.exists():
        RECENT_HISTORY_FILE.write_text("", encoding="utf-8")


def safe_path(user_path: str) -> Path:
    if not user_path or not user_path.strip():
        raise ValueError("路径不能为空")

    candidate = (WORKSPACE_ROOT / user_path).resolve()

    if candidate == WORKSPACE_ROOT:
        return candidate

    try:
        candidate.relative_to(WORKSPACE_ROOT)
    except ValueError as exc:
        raise ValueError("不允许访问工作区之外的路径") from exc

    return candidate


def is_text_file(path: Path) -> bool:
    return path.suffix.lower() in TEXT_EXTENSIONS


def is_docx_file(path: Path) -> bool:
    return path.suffix.lower() == ".docx"


def sanitize_filename(name: str) -> str:
    name = unquote(name)
    name = re.sub(r'[\\/:*?"<>|]+', "_", name).strip()
    return name or "downloaded_file"


def split_paragraphs(content: str) -> List[str]:
    text = content.replace("\r\n", "\n").strip()
    if not text:
        return [""]
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


def clear_docx_content(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        body.remove(child)


def list_dir(path: str = ".") -> Dict[str, Any]:
    target = safe_path(path)

    if not target.exists():
        return {"ok": False, "error": f"目录不存在: {target}"}
    if not target.is_dir():
        return {"ok": False, "error": f"目标不是目录: {target}"}

    items = []
    for item in sorted(target.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower())):
        items.append({
            "name": item.name,
            "type": "dir" if item.is_dir() else "file",
            "path": str(item.relative_to(WORKSPACE_ROOT))
        })

    return {
        "ok": True,
        "path": str(target.relative_to(WORKSPACE_ROOT)) if target != WORKSPACE_ROOT else ".",
        "items": items
    }


def create_folder(path: str) -> Dict[str, Any]:
    target = safe_path(path)
    target.mkdir(parents=True, exist_ok=True)
    return {
        "ok": True,
        "message": "文件夹创建成功",
        "path": str(target.relative_to(WORKSPACE_ROOT))
    }


def create_file(path: str, content: str = "") -> Dict[str, Any]:
    target = safe_path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    if target.exists() and target.is_dir():
        return {"ok": False, "error": "目标路径已存在，但它是一个文件夹，不是文件"}

    try:
        if is_docx_file(target):
            doc = Document()
            paragraphs = split_paragraphs(content)
            for para in paragraphs:
                doc.add_paragraph(para)
            doc.save(str(target))
        else:
            target.write_text(content, encoding="utf-8")

        return {
            "ok": True,
            "message": "文件创建成功",
            "path": str(target.relative_to(WORKSPACE_ROOT))
        }
    except PermissionError:
        return {
            "ok": False,
            "error": f"没有权限写入文件：{target}。常见原因是该文件正被 Word/WPS 打开，请先关闭后再试。"
        }
    except Exception as e:
        return {"ok": False, "error": f"{type(e).__name__}: {str(e)}"}


def read_file(path: str) -> Dict[str, Any]:
    target = safe_path(path)

    if not target.exists():
        return {"ok": False, "error": f"文件不存在: {target}"}
    if not target.is_file():
        return {"ok": False, "error": f"目标不是文件: {target}"}

    try:
        if is_docx_file(target):
            doc = Document(str(target))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            content = "\n".join(paragraphs)
            return {
                "ok": True,
                "path": str(target.relative_to(WORKSPACE_ROOT)),
                "content": content,
                "file_type": "docx"
            }

        if is_text_file(target):
            content = target.read_text(encoding="utf-8")
            return {
                "ok": True,
                "path": str(target.relative_to(WORKSPACE_ROOT)),
                "content": content,
                "file_type": "text"
            }

        return {
            "ok": False,
            "error": f"暂不支持直接读取这种文件类型：{target.suffix}。目前支持常见文本文件和 .docx"
        }
    except PermissionError:
        return {
            "ok": False,
            "error": f"没有权限读取文件：{target}。如果文件正被其他程序占用，请先关闭后再试。"
        }
    except Exception as e:
        return {"ok": False, "error": f"{type(e).__name__}: {str(e)}"}


def write_docx_file(target: Path, content: str, mode: str) -> Dict[str, Any]:
    try:
        if target.exists():
            doc = Document(str(target))
        else:
            doc = Document()

        if mode == "overwrite":
            clear_docx_content(doc)

        paragraphs = split_paragraphs(content)
        for para in paragraphs:
            doc.add_paragraph(para)

        doc.save(str(target))
        return {
            "ok": True,
            "message": "Word 文档写入成功",
            "path": str(target.relative_to(WORKSPACE_ROOT)),
            "mode": mode
        }
    except PermissionError:
        return {
            "ok": False,
            "error": f"没有权限写入文件：{target}。最常见原因是该 Word 文件正被 Word/WPS 打开，请先关闭后再试。"
        }
    except Exception as e:
        return {"ok": False, "error": f"{type(e).__name__}: {str(e)}"}


def write_file(path: str, content: str, mode: str = "overwrite") -> Dict[str, Any]:
    target = safe_path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    if target.exists() and target.is_dir():
        return {"ok": False, "error": "目标路径是文件夹，不能写入成文件"}

    try:
        if is_docx_file(target):
            return write_docx_file(target, content, mode)

        if mode == "append":
            with target.open("a", encoding="utf-8") as f:
                f.write(content)
        else:
            target.write_text(content, encoding="utf-8")

        return {
            "ok": True,
            "message": "写入成功",
            "path": str(target.relative_to(WORKSPACE_ROOT)),
            "mode": mode
        }
    except PermissionError:
        return {
            "ok": False,
            "error": f"没有权限写入文件：{target}。如果文件正被其他程序占用，请先关闭后再试。"
        }
    except Exception as e:
        return {"ok": False, "error": f"{type(e).__name__}: {str(e)}"}


def replace_in_file(path: str, old_text: str, new_text: str, count: int = -1) -> Dict[str, Any]:
    target = safe_path(path)

    if not target.exists():
        return {"ok": False, "error": "文件不存在"}
    if not target.is_file():
        return {"ok": False, "error": "目标不是文件"}

    try:
        if is_docx_file(target):
            doc = Document(str(target))
            full_text = "\n".join([p.text for p in doc.paragraphs])

            if old_text not in full_text:
                return {
                    "ok": False,
                    "error": "没有找到要替换的内容",
                    "path": str(target.relative_to(WORKSPACE_ROOT))
                }

            if count == -1:
                new_full_text = full_text.replace(old_text, new_text)
            else:
                new_full_text = full_text.replace(old_text, new_text, count)

            clear_docx_content(doc)
            for para in split_paragraphs(new_full_text):
                doc.add_paragraph(para)
            doc.save(str(target))

            return {
                "ok": True,
                "message": "Word 文档修改成功",
                "path": str(target.relative_to(WORKSPACE_ROOT))
            }

        content = target.read_text(encoding="utf-8")

        if old_text not in content:
            return {
                "ok": False,
                "error": "没有找到要替换的内容",
                "path": str(target.relative_to(WORKSPACE_ROOT))
            }

        new_content = content.replace(old_text, new_text) if count == -1 else content.replace(old_text, new_text, count)
        target.write_text(new_content, encoding="utf-8")

        return {
            "ok": True,
            "message": "修改成功",
            "path": str(target.relative_to(WORKSPACE_ROOT))
        }

    except PermissionError:
        return {
            "ok": False,
            "error": f"没有权限修改文件：{target}。如果这个 Word 文件正被打开，请先关闭后再试。"
        }
    except Exception as e:
        return {"ok": False, "error": f"{type(e).__name__}: {str(e)}"}


def guess_filename_from_response(response: requests.Response, url: str) -> str:
    cd = response.headers.get("Content-Disposition", "")

    match = re.search(r"filename\*=UTF-8''([^;]+)", cd)
    if match:
        return sanitize_filename(unquote(match.group(1)))

    match = re.search(r'filename="?([^"]+)"?', cd)
    if match:
        return sanitize_filename(match.group(1))

    parsed = urlparse(url)
    raw_name = Path(parsed.path).name
    return sanitize_filename(raw_name or "downloaded_file")


def download_url(
    url: str,
    save_path: Optional[str] = None,
    verify_ssl: bool = True
) -> Dict[str, Any]:
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https"):
        return {"ok": False, "error": "只允许下载 http 或 https 链接"}

    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/129.0.0.0 Safari/537.36 Edg/129.0.0.0"
        ),
        "Accept": "*/*",
        "Referer": "https://www.bing.com/"
    }

    try:
        response = requests.get(
            url,
            headers=headers,
            stream=True,
            timeout=120,
            allow_redirects=True,
            verify=verify_ssl
        )
        response.raise_for_status()

        if save_path:
            target = safe_path(save_path)
        else:
            filename = guess_filename_from_response(response, url)
            target = safe_path(f"downloads/{filename}")

        target.parent.mkdir(parents=True, exist_ok=True)

        with open(target, "wb") as f:
            for chunk in response.iter_content(chunk_size=8192):
                if chunk:
                    f.write(chunk)

        return {
            "ok": True,
            "message": "下载成功",
            "path": str(target.relative_to(WORKSPACE_ROOT)),
            "url": url,
            "content_type": response.headers.get("Content-Type", ""),
            "size_bytes": target.stat().st_size
        }

    except requests.exceptions.SSLError as e:
        return {
            "ok": False,
            "error": (
                f"SSL 证书校验失败：{str(e)}。"
                "如果您确认这个网址可信，可以把 verify_ssl 设为 false 再试一次。"
            )
        }
    except Exception as e:
        return {"ok": False, "error": f"{type(e).__name__}: {str(e)}"}


def build_edge_driver() -> webdriver.Edge:
    headless = os.getenv("BING_HEADLESS", "false").lower() == "true"

    options = Options()
    options.add_argument(
        "user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/129.0.0.0 Safari/537.36 Edg/129.0.0.0"
    )
    options.add_argument("--disable-blink-features=AutomationControlled")
    options.add_argument("--start-maximized")

    if headless:
        options.add_argument("--headless=new")

    driver = webdriver.Edge(options=options)
    driver.set_page_load_timeout(60)
    return driver


def web_search_bing_requests(query: str, count: int = 5) -> Dict[str, Any]:
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/129.0.0.0 Safari/537.36 Edg/129.0.0.0"
        )
    }

    try:
        url = f"https://www.bing.com/search?q={quote(query)}"
        resp = requests.get(url, headers=headers, timeout=30)
        resp.raise_for_status()

        soup = BeautifulSoup(resp.text, "lxml")
        results = []

        for li in soup.select("li.b_algo"):
            a = li.select_one("h2 a")
            p = li.select_one(".b_caption p")
            if not a:
                continue

            results.append({
                "title": a.get_text(" ", strip=True),
                "content": p.get_text(" ", strip=True) if p else "",
                "link": a.get("href", ""),
                "media": "Bing",
                "publish_date": ""
            })

            if len(results) >= max(1, min(count, 10)):
                break

        if not results:
            return {
                "ok": False,
                "error": "Bing 直连搜索没有解析到结果。可能是页面结构变化，或者当前网络环境限制了访问。"
            }

        return {
            "ok": True,
            "engine": "bing_requests_web",
            "query": query,
            "results": results
        }
    except Exception as e:
        return {"ok": False, "error": f"{type(e).__name__}: {str(e)}"}


def web_search_bing_selenium(query: str, count: int = 5) -> Dict[str, Any]:
    driver = None
    try:
        driver = build_edge_driver()
        url = f"https://www.bing.com/search?q={quote(query)}"
        driver.get(url)

        WebDriverWait(driver, 20).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "li.b_algo"))
        )

        items = driver.find_elements(By.CSS_SELECTOR, "li.b_algo")
        results = []

        for item in items:
            try:
                title_el = item.find_element(By.CSS_SELECTOR, "h2")
                link_el = item.find_element(By.CSS_SELECTOR, "h2 a")

                snippet = ""
                try:
                    snippet_el = item.find_element(By.CSS_SELECTOR, ".b_caption p")
                    snippet = snippet_el.text.strip()
                except Exception:
                    pass

                title = title_el.text.strip()
                link = link_el.get_attribute("href")

                if title and link:
                    results.append({
                        "title": title,
                        "content": snippet,
                        "link": link,
                        "media": "Bing",
                        "publish_date": ""
                    })

                if len(results) >= max(1, min(count, 10)):
                    break
            except Exception:
                continue

        if not results:
            return {
                "ok": False,
                "error": "Bing Selenium 搜索没有解析到结果。"
            }

        return {
            "ok": True,
            "engine": "bing_selenium_web",
            "query": query,
            "results": results
        }
    except Exception as e:
        return {"ok": False, "error": f"{type(e).__name__}: {str(e)}"}
    finally:
        if driver:
            driver.quit()


def web_search_bing(query: str, count: int = 5) -> Dict[str, Any]:
    first = web_search_bing_requests(query, count)
    if first.get("ok"):
        return first

    second = web_search_bing_selenium(query, count)
    if second.get("ok"):
        return second

    return {
        "ok": False,
        "error": (
            "Bing 搜索失败。"
            f"直连结果：{first.get('error', '未知错误')}；"
            f"Selenium 结果：{second.get('error', '未知错误')}"
        )
    }


def search_news_bing(query: str, count: int = 5) -> Dict[str, Any]:
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/129.0.0.0 Safari/537.36 Edg/129.0.0.0"
        )
    }

    try:
        url = f"https://www.bing.com/news/search?q={quote(query)}"
        resp = requests.get(url, headers=headers, timeout=30)
        resp.raise_for_status()

        soup = BeautifulSoup(resp.text, "lxml")
        results = []

        # 多给几套选择器，尽量兼容页面变化
        cards = soup.select("div.news-card, div.t_s, div.newsitem, div.card, div.caption")
        for card in cards:
            try:
                a = card.select_one("a.title") or card.select_one("a")
                if not a:
                    continue

                title = a.get_text(" ", strip=True)
                link = a.get("href", "")

                snippet_el = (
                    card.select_one(".snippet")
                    or card.select_one(".source")
                    or card.select_one("div")
                )
                snippet = snippet_el.get_text(" ", strip=True) if snippet_el else ""

                source_el = card.select_one(".source")
                source = source_el.get_text(" ", strip=True) if source_el else ""

                if title and link:
                    results.append({
                        "title": title,
                        "content": snippet,
                        "link": link,
                        "source": source
                    })

                if len(results) >= max(1, min(count, 10)):
                    break
            except Exception:
                continue

        if not results:
            fallback = web_search_bing(f"{query} 最新新闻", count=count)
            if fallback.get("ok"):
                return {
                    "ok": True,
                    "engine": "bing_news_fallback_web",
                    "query": query,
                    "results": fallback.get("results", [])
                }
            return {"ok": False, "error": "Bing News 没有解析到结果。"}

        return {
            "ok": True,
            "engine": "bing_news",
            "query": query,
            "results": results
        }

    except Exception as e:
        fallback = web_search_bing(f"{query} 最新新闻", count=count)
        if fallback.get("ok"):
            return {
                "ok": True,
                "engine": "bing_news_fallback_web",
                "query": query,
                "results": fallback.get("results", [])
            }
        return {"ok": False, "error": f"{type(e).__name__}: {str(e)}"}


def image_search_bing(query: str, count: int = 10) -> Dict[str, Any]:
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/129.0.0.0 Safari/537.36 Edg/129.0.0.0"
        )
    }

    # 先用 requests 直接抓图
    try:
        url = f"https://www.bing.com/images/search?q={quote(query)}"
        resp = requests.get(url, headers=headers, timeout=30)
        resp.raise_for_status()

        soup = BeautifulSoup(resp.text, "lxml")
        results = []
        seen = set()

        for a in soup.select("a.iusc"):
            try:
                meta_str = a.get("m")
                if not meta_str:
                    continue

                meta = json.loads(meta_str)
                img_url = meta.get("murl") or meta.get("turl")
                thumb_url = meta.get("turl", "")
                title = meta.get("t", "") or ""

                if img_url and img_url.startswith("http") and img_url not in seen:
                    seen.add(img_url)
                    results.append({
                        "url": img_url,
                        "thumbnail": thumb_url,
                        "title": title
                    })

                if len(results) >= max(1, min(count, 50)):
                    break
            except Exception:
                continue

        if results:
            return {
                "ok": True,
                "engine": "bing_images_requests",
                "query": query,
                "results": results[:count]
            }
    except Exception:
        pass

    # 再用 Selenium 兜底
    driver = None
    try:
        driver = build_edge_driver()
        url = f"https://www.bing.com/images/search?q={quote(query)}"
        driver.get(url)

        WebDriverWait(driver, 20).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "a.iusc"))
        )

        results = []
        seen = set()

        for _ in range(5):
            cards = driver.find_elements(By.CSS_SELECTOR, "a.iusc")
            for card in cards:
                try:
                    meta_str = card.get_attribute("m")
                    if not meta_str:
                        continue

                    meta = json.loads(meta_str)
                    img_url = meta.get("murl") or meta.get("turl")
                    thumb_url = meta.get("turl", "")
                    title = meta.get("t", "") or ""

                    if img_url and img_url.startswith("http") and img_url not in seen:
                        seen.add(img_url)
                        results.append({
                            "url": img_url,
                            "thumbnail": thumb_url,
                            "title": title
                        })

                    if len(results) >= max(1, min(count, 50)):
                        break
                except Exception:
                    continue

            if len(results) >= max(1, min(count, 50)):
                break

            driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
            time.sleep(2)

        if results:
            return {
                "ok": True,
                "engine": "bing_images_selenium",
                "query": query,
                "results": results[:count]
            }

        return {
            "ok": False,
            "error": "Bing 图片搜索没有抓到可用图片链接。"
        }

    except Exception as e:
        return {"ok": False, "error": f"{type(e).__name__}: {str(e)}"}
    finally:
        if driver:
            driver.quit()


def download_bing_image(
    query: str,
    save_path: str,
    index: int = 1,
    verify_ssl: bool = True
) -> Dict[str, Any]:
    search_result = image_search_bing(query=query, count=max(index, 5))
    if not search_result.get("ok"):
        return {
            "ok": False,
            "error": f"图片搜索失败：{search_result.get('error', '未知错误')}"
        }

    results = search_result.get("results", [])
    if not results or len(results) < index:
        return {
            "ok": False,
            "error": f"搜索结果不足，无法下载第 {index} 张图片。"
        }

    img_url = results[index - 1]["url"]
    download_result = download_url(
        url=img_url,
        save_path=save_path,
        verify_ssl=verify_ssl
    )

    if download_result.get("ok"):
        download_result["source_image_url"] = img_url

    return download_result


def get_current_datetime() -> Dict[str, Any]:
    now = datetime.now()
    return {
        "ok": True,
        "date": now.strftime("%Y-%m-%d"),
        "time": now.strftime("%H:%M:%S"),
        "datetime": now.strftime("%Y-%m-%d %H:%M:%S"),
        "weekday": now.strftime("%A")
    }


def load_memory_summary() -> str:
    if MEMORY_SUMMARY_FILE.exists():
        return MEMORY_SUMMARY_FILE.read_text(encoding="utf-8").strip()
    return ""


def load_recent_history() -> str:
    if RECENT_HISTORY_FILE.exists():
        return RECENT_HISTORY_FILE.read_text(encoding="utf-8").strip()
    return ""


def append_recent_history(user_text: str, assistant_text: str) -> None:
    old_text = load_recent_history()
    block = (
        "[ROUND]\n"
        f"USER: {user_text}\n"
        f"ASSISTANT: {assistant_text}\n"
        "[/ROUND]\n"
    )

    merged = (old_text + "\n" + block).strip()
    rounds = re.findall(r"\[ROUND\]\n.*?\n\[/ROUND\]", merged, flags=re.DOTALL)
    rounds = rounds[-MAX_RECENT_ROUNDS:]

    RECENT_HISTORY_FILE.write_text("\n".join(rounds), encoding="utf-8")


def extract_text_content(message: Dict[str, Any]) -> str:
    content = message.get("content", "")

    if isinstance(content, str):
        return content

    if isinstance(content, list):
        parts: List[str] = []
        for item in content:
            if isinstance(item, dict) and item.get("type") == "text":
                parts.append(item.get("text", ""))
        return "\n".join(parts).strip()

    return ""


def call_llm(messages: List[Dict[str, Any]], tools: Optional[List[Dict[str, Any]]] = None) -> Dict[str, Any]:
    if not LLM_API_KEY:
        raise RuntimeError("未在 .env 中找到 ZHIPU_API_KEY，请先配置。")

    headers = {
        "Authorization": f"Bearer {LLM_API_KEY}",
        "Content-Type": "application/json",
    }

    payload: Dict[str, Any] = {
        "model": LLM_MODEL,
        "messages": messages,
        "stream": False,
        "temperature": 0.3,
    }

    if tools:
        payload["tools"] = tools
        payload["tool_choice"] = "auto"

    response = requests.post(LLM_API_URL, headers=headers, json=payload, timeout=120)
    response.raise_for_status()
    return response.json()


def refresh_memory_summary() -> None:
    recent = load_recent_history()
    old_summary = load_memory_summary()

    if not recent.strip():
        return

    messages = [
        {
            "role": "system",
            "content": (
                "请把下面的用户历史对话压缩成一份简洁、稳定、有用的长期记忆摘要。"
                "重点保留：用户偏好、长期任务背景、反复提到的文件路径、常用工作方式。"
                "不要保留无意义寒暄。输出中文纯文本。"
            )
        },
        {
            "role": "user",
            "content": (
                f"旧摘要：\n{old_summary}\n\n"
                f"最近对话：\n{recent}\n\n"
                "请输出新的整合摘要。"
            )
        }
    ]

    try:
        result = call_llm(messages, tools=None)
        choices = result.get("choices") or []
        if choices:
            msg = choices[0].get("message") or {}
            summary = extract_text_content(msg).strip()
            if summary:
                MEMORY_SUMMARY_FILE.write_text(summary, encoding="utf-8")
    except Exception:
        pass


LOCAL_TOOLS = {
    "list_dir": list_dir,
    "create_folder": create_folder,
    "create_file": create_file,
    "read_file": read_file,
    "write_file": write_file,
    "replace_in_file": replace_in_file,
    "web_search_bing": web_search_bing,
    "search_news_bing": search_news_bing,
    "image_search_bing": image_search_bing,
    "download_url": download_url,
    "download_bing_image": download_bing_image,
    "get_current_datetime": get_current_datetime,
}

TOOL_SCHEMAS = [
    {
        "type": "function",
        "function": {
            "name": "list_dir",
            "description": "读取某个目录下的文件和文件夹列表",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "相对工作区的目录路径"}
                }
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "create_folder",
            "description": "创建文件夹，可自动创建多级目录",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "相对工作区的文件夹路径"}
                },
                "required": ["path"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "create_file",
            "description": "创建文件。支持常见文本文件和 .docx 文件",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "相对工作区的文件路径"},
                    "content": {"type": "string", "description": "初始内容", "default": ""}
                },
                "required": ["path"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "read_file",
            "description": "读取文件内容。当前支持常见文本文件和 .docx 文件",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "相对工作区的文件路径"}
                },
                "required": ["path"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "write_file",
            "description": "向文件写入内容。支持常见文本文件和 .docx 文件",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "相对工作区的文件路径"},
                    "content": {"type": "string", "description": "要写入的内容"},
                    "mode": {
                        "type": "string",
                        "enum": ["overwrite", "append"],
                        "description": "overwrite=覆盖写入，append=追加写入",
                        "default": "overwrite"
                    }
                },
                "required": ["path", "content"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "replace_in_file",
            "description": "修改文件内容，把旧文本替换成新文本。支持常见文本文件和 .docx 文件",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "相对工作区的文件路径"},
                    "old_text": {"type": "string", "description": "旧内容"},
                    "new_text": {"type": "string", "description": "新内容"},
                    "count": {"type": "integer", "description": "替换次数，-1 表示全部替换", "default": -1}
                },
                "required": ["path", "old_text", "new_text"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "web_search_bing",
            "description": "使用 Bing 网页搜索查询资料，返回标题、摘要和链接。会优先尝试普通网页请求，失败时再尝试浏览器自动化。",
            "parameters": {
                "type": "object",
                "properties": {
                    "query": {"type": "string", "description": "搜索关键词"},
                    "count": {"type": "integer", "description": "最多返回多少条结果，建议 1 到 10", "default": 5}
                },
                "required": ["query"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "search_news_bing",
            "description": "使用 Bing News 搜索最近新闻，并返回标题、摘要和链接。",
            "parameters": {
                "type": "object",
                "properties": {
                    "query": {"type": "string", "description": "新闻搜索关键词"},
                    "count": {"type": "integer", "description": "返回结果数量", "default": 5}
                },
                "required": ["query"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "image_search_bing",
            "description": "使用 Bing 图片搜索，返回图片链接列表。",
            "parameters": {
                "type": "object",
                "properties": {
                    "query": {"type": "string", "description": "图片搜索关键词，例如：兔子的图片"},
                    "count": {"type": "integer", "description": "最多返回多少条图片链接，建议 1 到 20", "default": 10}
                },
                "required": ["query"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "download_url",
            "description": "从网上下载图片、文档或其他文件到本地工作区。",
            "parameters": {
                "type": "object",
                "properties": {
                    "url": {"type": "string", "description": "要下载的 http 或 https 链接"},
                    "save_path": {"type": "string", "description": "保存到工作区中的相对路径，可不填"},
                    "verify_ssl": {
                        "type": "boolean",
                        "description": "是否校验 SSL 证书。默认为 true；如果证书链异常但网址可信，可设为 false。",
                        "default": True
                    }
                },
                "required": ["url"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "download_bing_image",
            "description": "先搜索图片，再把搜索结果中的第 index 张图片下载到本地。",
            "parameters": {
                "type": "object",
                "properties": {
                    "query": {"type": "string", "description": "图片搜索关键词"},
                    "save_path": {"type": "string", "description": "保存到工作区中的相对路径"},
                    "index": {"type": "integer", "description": "第几张图，从1开始", "default": 1},
                    "verify_ssl": {"type": "boolean", "description": "是否校验SSL证书", "default": True}
                },
                "required": ["query", "save_path"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "get_current_datetime",
            "description": "获取当前系统本地日期、时间和完整时间字符串。",
            "parameters": {
                "type": "object",
                "properties": {}
            }
        }
    }
]


def execute_tool_call(tc: Dict[str, Any]) -> Dict[str, Any]:
    func = tc.get("function") or {}
    tool_name = func.get("name")
    arg_str = func.get("arguments") or "{}"

    try:
        args = json.loads(arg_str) if isinstance(arg_str, str) else arg_str
    except Exception:
        args = {}

    if tool_name not in LOCAL_TOOLS:
        tool_result = {"ok": False, "error": f"未知工具: {tool_name}"}
    else:
        try:
            tool_result = LOCAL_TOOLS[tool_name](**args)
        except Exception as e:
            tool_result = {"ok": False, "error": f"{type(e).__name__}: {str(e)}"}

    return {
        "tool_name": tool_name,
        "args": args,
        "result": tool_result,
        "tool_call_id": tc.get("id")
    }


def run_agent_turn(conversation: List[Dict[str, Any]], user_text: str) -> Dict[str, Any]:
    conversation.append({"role": "user", "content": user_text})
    executed_tools: List[Dict[str, Any]] = []

    for _step in range(MAX_AGENT_STEPS):
        resp = call_llm(conversation, TOOL_SCHEMAS)
        choices = resp.get("choices") or []
        if not choices:
            return {
                "reply": "模型没有返回有效结果。",
                "tool_results": executed_tools
            }

        message = choices[0].get("message") or {}
        tool_calls = message.get("tool_calls") or []
        assistant_text = extract_text_content(message)

        assistant_record: Dict[str, Any] = {
            "role": "assistant",
            "content": assistant_text,
        }
        if tool_calls:
            assistant_record["tool_calls"] = tool_calls
        conversation.append(assistant_record)

        if not tool_calls:
            final_text = assistant_text or "已经执行完成，但模型没有返回文字说明。"
            return {
                "reply": final_text,
                "tool_results": executed_tools
            }

        for tc in tool_calls:
            result = execute_tool_call(tc)
            executed_tools.append(result)

            conversation.append({
                "role": "tool",
                "tool_call_id": result["tool_call_id"],
                "content": json.dumps(result["result"], ensure_ascii=False)
            })

    return {
        "reply": f"任务执行到第 {MAX_AGENT_STEPS} 步后停止了，可能是模型陷入重复调用。请把需求说得再具体一点，我再继续帮您处理。",
        "tool_results": executed_tools
    }


def print_tool_results(tool_results: List[Dict[str, Any]]) -> None:
    if not tool_results:
        return

    print("\n[工具执行记录]")
    for idx, item in enumerate(tool_results, start=1):
        print(f"{idx}. 工具名: {item['tool_name']}")
        print(f"   参数: {json.dumps(item['args'], ensure_ascii=False)}")
        print(f"   结果: {json.dumps(item['result'], ensure_ascii=False)}")
    print()


def main() -> None:
    ensure_workspace()

    print("======================================")
    print("本地文件工作助手 已启动")
    print(f"工作区: {WORKSPACE_ROOT}")
    print("输入内容后按回车即可对话")
    print("输入“结束对话”即可退出")
    print("======================================\n")

    print("联网工具说明：")
    print("1. 普通网页搜索优先使用 requests 直连 Bing。")
    print("2. 新闻搜索优先使用 Bing News。")
    print("3. 图片搜索先尝试 requests，失败后再尝试 Edge。")
    print("4. 图片下载优先使用 download_bing_image。")
    print("5. 程序会把最近几轮对话和长期摘要写入工作区记忆文件。\n")

    memory_summary = load_memory_summary()
    recent_history = load_recent_history()

    conversation: List[Dict[str, Any]] = [
        {"role": "system", "content": LLM_SYSTEM_PROMPT}
    ]

    if memory_summary:
        conversation.append({
            "role": "system",
            "content": f"以下是该用户的长期记忆摘要，请在回答时参考：\n{memory_summary}"
        })

    if recent_history:
        conversation.append({
            "role": "system",
            "content": f"以下是最近几轮对话记录，请在当前会话中参考：\n{recent_history}"
        })

    while True:
        try:
            user_text = input("你：").strip()
        except (EOFError, KeyboardInterrupt):
            print("\n助手：本次对话已结束，再见。")
            break

        if not user_text:
            continue

        if user_text == "结束对话":
            print("助手：好的，本次对话已经结束。欢迎下次再来。")
            break

        try:
            result = run_agent_turn(conversation, user_text)
            print(f"\n助手：{result['reply']}\n")
            print_tool_results(result.get("tool_results", []))

            append_recent_history(user_text, result["reply"])
            refresh_memory_summary()

        except Exception as e:
            print(f"\n助手：程序运行时发生错误：{type(e).__name__}: {str(e)}\n")


if __name__ == "__main__":
    main()